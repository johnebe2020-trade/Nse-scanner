import os, time, csv
from datetime import datetime, timezone, timedelta
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote
import numpy as np
import requests
from openpyxl import load_workbook

# ============ EDIT HERE ============
STOCK_FILE = "stocks.xlsx"     # your Excel file name
EXIT_LOOKBACK = 1              # show EXIT rows if exit happened in last N bars
WORKERS = 8
SYMBOL_FIX = {}                # rename: {"OLDNAME": "YAHOONAME"}
REMOVE = set()                 # drop: {"EBIX", "COMPEAU"}
ST_PERIOD, ST_FACTOR = 10, 2.0
RSI_LEN, RSI_MIN = 14, 70
BB_LEN, BB_MULT = 20, 2.0
# ===================================

IST = timezone(timedelta(hours=5, minutes=30))
UA = {"User-Agent": "Mozilla/5.0"}


def load_stocks():
    ws = load_workbook(STOCK_FILE, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    pairs = [i for i, v in enumerate(rows[0]) if str(v).strip().upper() == "STOCK"]
    out = {}
    for r in rows[1:]:
        for i in pairs:
            if i + 1 >= len(r) or not r[i]:
                continue
            s = str(r[i]).upper().replace("NSE:", "").replace("NSE ", "").replace(" ", "")
            s = SYMBOL_FIX.get(s, s)
            if not s or s in REMOVE:
                continue
            sec = str(r[i + 1] or "?").strip().upper()
            out.setdefault(s, [])
            if sec not in out[s]:
                out[s].append(sec)
    return out


def fetch(sym):
    url = "https://query1.finance.yahoo.com/v8/finance/chart/" + quote(sym + ".NS", safe="")
    err = ("E3", "FETCH ERROR")
    for a in range(3):
        try:
            r = requests.get(url, params={"range": "1y", "interval": "1d"}, headers=UA, timeout=20)
            if r.status_code == 404:
                return None, ("E1", "INVALID SYMBOL")
            if r.status_code == 200:
                res = r.json()["chart"]["result"]
                if not res:
                    return None, ("E1", "INVALID SYMBOL")
                res = res[0]
                q = res["indicators"]["quote"][0]
                rows = [(t, h, l, c) for t, h, l, c in
                        zip(res.get("timestamp") or [], q["high"], q["low"], q["close"])
                        if h is not None and l is not None and c is not None]
                return rows, None
            err = ("E3", "HTTP " + str(r.status_code))
        except Exception as e:
            err = ("E3", "FETCH " + type(e).__name__)
        time.sleep(1.5 * (a + 1))
    return None, err


def rma(x, n):
    out = np.full(len(x), np.nan)
    if len(x) < n:
        return out
    out[n - 1] = np.mean(x[:n])
    for i in range(n, len(x)):
        out[i] = (out[i - 1] * (n - 1) + x[i]) / n
    return out


def rsi(c, n):
    d = np.diff(c)
    up, dn = rma(np.maximum(d, 0), n), rma(np.maximum(-d, 0), n)
    with np.errstate(divide="ignore", invalid="ignore"):
        r = 100 - 100 / (1 + up / dn)
    return np.concatenate([[np.nan], r])


def supertrend(h, l, c, n, f):
    pc = np.concatenate([[c[0]], c[:-1]])
    tr = np.maximum(h - l, np.maximum(abs(h - pc), abs(l - pc)))
    tr[0] = h[0] - l[0]
    atr = rma(tr, n)
    hl2 = (h + l) / 2
    up, lo = hl2 + f * atr, hl2 - f * atr
    st = np.full(len(c), np.nan)
    d = np.ones(len(c))
    for i in range(len(c)):
        if np.isnan(atr[i]):
            continue
        if i > 0 and not np.isnan(atr[i - 1]):
            if not (lo[i] > lo[i - 1] or c[i - 1] < lo[i - 1]):
                lo[i] = lo[i - 1]
            if not (up[i] < up[i - 1] or c[i - 1] > up[i - 1]):
                up[i] = up[i - 1]
            if st[i - 1] == up[i - 1]:
                d[i] = -1 if c[i] > up[i] else 1
            else:
                d[i] = 1 if c[i] < lo[i] else -1
        st[i] = lo[i] if d[i] == -1 else up[i]
    return st


def fdate(t):
    return datetime.fromtimestamp(t, IST).strftime("%d-%b-%y")


def analyse(sym):
    sec = "/".join(STOCKS[sym])
    rows, err = fetch(sym)
    if not err:
        if len(rows) < 40:
            err = ("E2", "NO DATA")
        elif (datetime.now(IST) - datetime.fromtimestamp(rows[-1][0], IST)).days > 7:
            err = ("E4", "STALE " + fdate(rows[-1][0]))
    if err:
        return dict(status="ERR", sym=sym, sector=sec, rsi="-", close="-",
                    when="-", note=err[0] + " " + err[1])

    ts = [r[0] for r in rows]
    h, l, c = (np.array([r[i] for r in rows], float) for i in (1, 2, 3))
    n = len(c)

    st = supertrend(h, l, c, ST_PERIOD, ST_FACTOR)
    rs = rsi(c, RSI_LEN)
    pp = (h + l + c) / 3
    r1 = np.concatenate([[np.nan], 2 * pp[:-1] - l[:-1]])   # previous day's traditional R1
    ub = np.full(n, np.nan)
    w = np.lib.stride_tricks.sliding_window_view(c, BB_LEN)
    ub[BB_LEN - 1:] = w.mean(1) + BB_MULT * w.std(1)

    c1, c2, c3, c4 = c > st, rs > RSI_MIN, c > r1, c > ub
    valid = ~(np.isnan(st) | np.isnan(rs) | np.isnan(r1) | np.isnan(ub))
    ok = valid & c1 & c2 & c3 & c4

    base = dict(sym=sym, sector=sec, rsi=f"{rs[-1]:.1f}", close=f"{c[-1]:.2f}")
    if ok[-1]:
        j = n - 1
        while j > 0 and ok[j - 1]:
            j -= 1
        return dict(base, status="ACTIVE", when=fdate(ts[j]), note="4/4 OK")

    idx = np.where(ok)[0]
    if len(idx) and (n - 1 - idx[-1]) <= EXIT_LOOKBACK:
        e = idx[-1] + 1
        failed = [nm for nm, cond in (("ST", c1), ("RSI", c2), ("R1", c3), ("BB", c4)) if not cond[e]]
        return dict(base, status="EXIT", when=fdate(ts[e]), note="FAIL " + " ".join(failed))
    return None


def sort_rows(res):
    order = {"ACTIVE": 0, "EXIT": 1, "ERR": 2}
    return sorted(res, key=lambda x: (order[x["status"]], x["sym"]))


if __name__ == "__main__":
    STOCKS = load_stocks()
    syms = list(STOCKS)
    print("Scanning", len(syms), "unique symbols")
    with ThreadPoolExecutor(WORKERS) as ex:
        out = sort_rows([r for r in ex.map(analyse, syms) if r])

    na_, nx, ne = (sum(r["status"] == s for r in out) for s in ("ACTIVE", "EXIT", "ERR"))
    summary = (f"4of4 Scan {datetime.now(IST):%d-%b-%y %H:%M} IST | ACTIVE {na_} | "
               f"EXIT {nx} | ERR {ne} | scanned {len(syms)}")
    print(summary)

    # CSV
    with open("results_4of4.csv", "w", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["status", "symbol", "sector", "rsi", "close", "date", "note"])
        for r in out:
            wr.writerow([r["status"], r["sym"], r["sector"], r["rsi"], r["close"], r["when"], r["note"]])

    # Markdown table (renders on GitHub)
    md = [f"**{summary}**", "",
          "| Status | Symbol | Sector | RSI | Close | Date | Note |",
          "|---|---|---|---|---|---|---|"]
    for r in out:
        md.append(f"| {r['status']} | {r['sym']} | {r['sector']} | {r['rsi']} | "
                  f"{r['close']} | {r['when']} | {r['note']} |")
    text = "\n".join(md)
    with open("results_4of4.md", "w") as f:
        f.write(text)
    print(text)

    # Actions run summary
    summ = os.getenv("GITHUB_STEP_SUMMARY")
    if summ:
        with open(summ, "a") as f:
            f.write(text)
